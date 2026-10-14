from __future__ import annotations

import csv
import json
import re
import sys
import zipfile
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PINN_ROOT = ROOT / "PINN_v2"
sys.path.insert(0, str(PINN_ROOT))

from pinn_v2.cases.definition import CaseDefinition, SCHEMA_VERSION  # noqa: E402


NAVY = "17365D"
RED = "8B1A1A"
LIGHT_BLUE = "DCE6F1"
MID_BLUE = "9FBAD0"
GREEN = "5A7D71"
GOLD = "E9B44C"
PALE_GOLD = "F8EFD5"
WHITE = "FFFFFF"
INK = "243746"
LIGHT_GRAY = "F2F2F2"
GRID = "CBD3DA"

AUTHORS = "Luis Enrique Koc Góngora y Herbert Antonio Meléndez García"
ADVISER = "Samuel Oporto"
THESIS_TITLE = (
    "Diseño y evaluación de una red neuronal informada por física para estimar "
    "la temperatura y la ampacidad de cables eléctricos enterrados en entornos heterogéneos"
)

TASKS = [
    (1, "Inventariar plan vigente, Guía 1 y observaciones previas", 1, 4, "-", "Base documental disponible; fechas semanales no trazables"),
    (2, "Verificar estructura profesionalizante, datos generales y APA", 1, 5, "1", "Base documental disponible; revisión continúa"),
    (3, "Revisar coherencia de diagnóstico, preguntas, objetivos y metodología", 2, 6, "1", "Base parcial; revisión continúa"),
    (4, "Consolidar inventario de datos, campos y procedencia", 2, 7, "1", "Inventario previo disponible; auditoría continúa"),
    (5, "Buscar proyectos COES y activos MINEM de cable subterráneo, cualquier tensión", 4, 7, "4", "En curso al 10-oct-2026"),
    (6, "Contrastar Arias, otros casos peruanos y CIGRE TB 880/TB 963", 5, 8, "5", "Programado"),
    (7, "Reconciliar datos y estados de procedencia con PINN_v2", 6, 9, "4, 6", "Programado"),
    (8, "Redactar informe de datos y construir libro Excel navegable", 7, 10, "4, 5, 6", "Programado"),
    (9, "Preparar paquete JSON, candidatos faltantes y manifiesto", 6, 9, "7", "Programado"),
    (10, "Evaluar revistas y redactar propuesta justificada", 9, 11, "6", "Programado"),
    (11, "Incorporar comentarios del asesor y corregir el plan", 10, 12, "1-10", "Programado"),
    (12, "Compilar, validar consistencia y cerrar los entregables", 12, 12, "1-11", "Programado"),
]

MINEM_ROWS = [
    {"id": "minem-2024-l2005-2006", "codigo": "L-2005/L-2006", "salida": "Chavarría", "llegada": "Barsi", "km": 1.00, "kv": 220, "mva": 335, "conductor": "Cu", "mm2": 1200, "linea_txt": 31484},
    {"id": "minem-2024-l2114", "codigo": "L-2114", "salida": "Mirador", "llegada": "Malvinas", "km": 4.08, "kv": 220, "mva": None, "conductor": "Cu", "mm2": 1200, "linea_txt": 31503},
    {"id": "minem-2024-l2115", "codigo": "L-2115", "salida": "Mirador", "llegada": "Malvinas", "km": 4.08, "kv": 220, "mva": None, "conductor": "Cu", "mm2": 1200, "linea_txt": 31507},
    {"id": "minem-2024-l607-282", "codigo": "L-607", "salida": "Santa Rosa Nueva", "llegada": "Tacna", "km": 2.82, "kv": 60, "mva": 23, "conductor": "Cu", "mm2": 120, "linea_txt": 31513},
    {"id": "minem-2024-l613", "codigo": "L-613", "salida": "Santa Rosa Antigua", "llegada": "Tacna", "km": 2.95, "kv": 60, "mva": 62, "conductor": "Cu", "mm2": 400, "linea_txt": 31519},
    {"id": "minem-2024-l615-a", "codigo": "L-615 (registro 1)", "salida": "Barsi", "llegada": "Pershing", "km": 1.04, "kv": 60, "mva": 78, "conductor": "Cu", "mm2": 500, "linea_txt": 31527},
    {"id": "minem-2024-l615-b", "codigo": "L-615 (registro 2)", "salida": "Barsi", "llegada": "Pershing", "km": 1.90, "kv": 60, "mva": 78, "conductor": "Cu", "mm2": 500, "linea_txt": 31531},
    {"id": "minem-2024-l621-622", "codigo": "L-621/L-622", "salida": "Chavarría", "llegada": "Mirones", "km": 1.89, "kv": 60, "mva": 56, "conductor": "Cu", "mm2": 500, "linea_txt": 31547},
    {"id": "minem-2024-l6349", "codigo": "L-6349", "salida": "Huarangal", "llegada": "Comas", "km": 7.84, "kv": 60, "mva": None, "conductor": "Cu", "mm2": 1200, "linea_txt": 31565},
    {"id": "minem-2024-l6556", "codigo": "L-6556", "salida": "Chavarría", "llegada": "UNI", "km": 5.00, "kv": 60, "mva": 67, "conductor": "Cu", "mm2": 500, "linea_txt": 31630},
    {"id": "minem-2024-l6740", "codigo": "L-6740", "salida": "Mirador", "llegada": "Mariátegui", "km": 3.63, "kv": 60, "mva": 90, "conductor": "Cu", "mm2": 500, "linea_txt": 31652},
]

COES_ROWS = [
    {
        "id": "coes-2023-chavarria-san-martin-220kv",
        "tipo": "Proyecto planificado",
        "codigo": "Enlace Subterráneo 220 kV Chavarría-San Martín",
        "salida": "Chavarría",
        "llegada": "San Martín (nueva SE)",
        "kv": 220,
        "km": None,
        "estado": "Incluido en el diagnóstico COES 2025-2034; ejecución no verificada",
        "localizador": "Informe COES/DP-01-2023, p. 69 impresa",
        "fuente": "Plan/datos_fuente/extractos_texto/COES_-_2023_-_Informe_de_diagnstico_de_las_condiciones_operativas_del_SEIN__periodo_2025--2034.txt",
        "datos": "Se mencionan enlace subterráneo, nueva SE GIS y acometidas subterráneas; no se publican geometría de cable/zanja ni propiedades térmicas del suelo en el extracto consultado.",
    },
    {
        "id": "coes-2023-manchay-pachacutec-220kv",
        "tipo": "Enlace citado como existente/operativo",
        "codigo": "L-2134 / L-2133",
        "salida": "Manchay",
        "llegada": "Pachacútec",
        "kv": 220,
        "km": None,
        "estado": "Diagnóstico COES describe el enlace como subterráneo; página impresa por confirmar",
        "localizador": "Sección de análisis operativo, texto extraído; línea de texto 8411",
        "fuente": "Plan/datos_fuente/extractos_texto/COES_-_2023_-_Informe_de_diagnstico_de_las_condiciones_operativas_del_SEIN__periodo_2025--2034.txt",
        "datos": "El COES menciona que el enlace subterráneo no cuenta con reactores de línea; se requiere cotejar el PDF y el estado actual del activo.",
    },
    {
        "id": "coes-2023-pachacutec-progreso-220kv",
        "tipo": "Enlace citado como existente/operativo",
        "codigo": "L-2135 / L-2136",
        "salida": "Pachacútec",
        "llegada": "Progreso",
        "kv": 220,
        "km": None,
        "estado": "Diagnóstico COES describe el enlace como subterráneo; página impresa por confirmar",
        "localizador": "Sección de análisis operativo, texto extraído; línea de texto 8411",
        "fuente": "Plan/datos_fuente/extractos_texto/COES_-_2023_-_Informe_de_diagnstico_de_las_condiciones_operativas_del_SEIN__periodo_2025--2034.txt",
        "datos": "El COES menciona que el enlace subterráneo no cuenta con reactores de línea; se requiere cotejar el PDF y el estado actual del activo.",
    },
]

DOI_FIXES = {
    "atoccsa2024": "10.3390/en17051023",
    "oclon2015": "10.1016/j.energy.2015.04.100",
    "kim2025": "10.1016/j.geothermics.2024.103151",
    "cigre2022_case0": "CIGRE TB 880",
    "cigre2025_case1": "CIGRE TB 963",
    "cigre2025_annulus": "CIGRE TB 963",
}

ACTIVE_MAP = {
    "kim2025": "kim_v5_summer",
    "aras2005": "aras_single; aras_flat",
    "cigre2025_annulus": "annulus",
    "mms2d_suite": "mms_constant; mms_high_contrast; mms_interface; mms_layered_y; mms_robin; mms_smooth_2d; mms_variable",
}


def read_prior_inventory() -> list[dict[str, str]]:
    path = ROOT / "Plan" / "envío_rev1" / "04_Datos_objeto_estudio" / "04_indice_instancias.csv"
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def read_minem_extracts() -> list[dict[str, object]]:
    path = ROOT / "Plan" / "datos_fuente" / "descargas" / "minem_anuario_2024.txt"
    rows = []
    for line_number, raw in enumerate(path.read_text(encoding="utf-8").split("\n"), start=1):
        if "Subterráneo" not in raw:
            continue
        match = re.search(
            r"(?P<km>[\d.,]+)\s+Subterráneo\s+(?P<kv>\d{2,3})\s+"
            r"(?P<mva>-|[\d.,]+)\s+(?P<conductor>.+?)\s+(?P<section>[\d.,]+)\s*$",
            raw,
            flags=re.IGNORECASE | re.DOTALL,
        )
        code_match = re.search(r"\bL-[A-Z0-9/-]+", raw, flags=re.IGNORECASE)
        fields = match.groupdict() if match else {}
        rows.append(
            {
                "id": f"minem2024-textrow-{line_number}",
                "linea_txt": line_number,
                "codigo": code_match.group(0) if code_match else "",
                "km": float(fields["km"].replace(",", ".")) if fields.get("km") else None,
                "kv": int(fields["kv"]) if fields.get("kv") else None,
                "mva": float(fields["mva"].replace(",", ".")) if fields.get("mva") and fields["mva"] != "-" else None,
                "conductor": " ".join(fields["conductor"].split()) if fields.get("conductor") else "",
                "mm2": float(fields["section"].replace(",", ".")) if fields.get("section") else None,
                "raw": raw.strip(),
                "parse_status": "Campos tabulares reconocidos" if match else "Solo extracto crudo; no parseado automáticamente",
            }
        )
    return rows


def load_active_cases() -> tuple[list[tuple[Path, CaseDefinition]], list[dict[str, object]]]:
    case_root = PINN_ROOT / "pinn_v2" / "cases" / "v2"
    loaded: list[tuple[Path, CaseDefinition]] = []
    for path in sorted(case_root.rglob("*.json")):
        definition = CaseDefinition.load(path)
        loaded.append((path, definition))
    status_rows = []
    for _, case in loaded:
        statuses = [entry.get("status") for entry in case.data_status.values()]
        status_rows.append(
            {
                "id": case.id,
                "tipo": case.physical.get("kind"),
                "estado_cargador": "Aceptado por CaseDefinition.load",
                "grupos_supuesto": sum(value == "supuesto" for value in statuses),
                "grupos_extraido": sum(value == "extraido" for value in statuses),
                "grupos_publicado": sum(value == "publicado" for value in statuses),
                "procedencia": "; ".join(f"{key}={value}" for key, value in case.data_status.items()),
            }
        )
    return loaded, status_rows


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def add_cover(workbook: Workbook, document_title: str, subtitle: str) -> None:
    sheet = workbook.create_sheet("Carátula", 0)
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H1")
    sheet["A1"] = "UNIVERSIDAD NACIONAL DE INGENIERÍA"
    sheet["A1"].font = Font(name="Arial", size=18, bold=True, color=NAVY)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:H2")
    sheet["A2"] = "UNIDAD DE POSGRADO DE LA FIIS  |  ESCUELA DE POSGRADO"
    sheet["A2"].font = Font(name="Arial", size=11, bold=True, color=RED)
    sheet["A2"].alignment = Alignment(horizontal="center")
    logo = ROOT / "Plan" / "imagenes" / "logo_UNI.png"
    if logo.exists():
        image = Image(str(logo))
        image.width = 126
        image.height = 126
        sheet.add_image(image, "D4")
    sheet.merge_cells("A11:H12")
    sheet["A11"] = document_title
    sheet["A11"].font = Font(name="Arial", size=16, bold=True, color=NAVY)
    sheet["A11"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.merge_cells("A14:H15")
    sheet["A14"] = THESIS_TITLE.upper()
    sheet["A14"].font = Font(name="Arial", size=10, bold=True, color=INK)
    sheet["A14"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.merge_cells("A17:H17")
    sheet["A17"] = "MAESTRÍA EN CIENCIAS CON MENCIÓN EN INTELIGENCIA ARTIFICIAL"
    sheet["A17"].font = Font(name="Arial", size=10, bold=True, color=INK)
    sheet["A17"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A19:H19")
    sheet["A19"] = f"TESISTAS: {AUTHORS}"
    sheet["A19"].font = Font(name="Arial", size=10, bold=True, color=INK)
    sheet["A19"].alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.merge_cells("A20:H20")
    sheet["A20"] = f"ASESOR: {ADVISER}"
    sheet["A20"].font = Font(name="Arial", size=10, color=INK)
    sheet["A20"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A22:H22")
    sheet["A22"] = subtitle
    sheet["A22"].font = Font(name="Arial", size=9, italic=True, color=INK)
    sheet["A22"].alignment = Alignment(horizontal="center", wrap_text=True)
    sheet.merge_cells("A24:H24")
    sheet["A24"] = "LIMA, PERÚ  |  2026"
    sheet["A24"].font = Font(name="Arial", size=10, color=INK)
    sheet["A24"].alignment = Alignment(horizontal="center")
    for row in range(1, 25):
        sheet.row_dimensions[row].height = 22
    sheet.row_dimensions[4].height = 24
    sheet.row_dimensions[11].height = 30
    sheet.row_dimensions[12].height = 28
    sheet.row_dimensions[14].height = 28
    sheet.row_dimensions[15].height = 24
    for col in "ABCDEFGH":
        sheet.column_dimensions[col].width = 15
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_area = "A1:H24"


def write_table(sheet, headers: list[str], rows: list[list[object]], table_name: str) -> None:
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = fill(NAVY)
        cell.font = Font(name="Arial", bold=True, color=WHITE, size=10)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=GRID))
    if rows:
        ref = f"A1:{sheet.cell(row=sheet.max_row, column=sheet.max_column).coordinate}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def style_workbook(workbook: Workbook) -> None:
    for sheet in workbook.worksheets:
        if sheet.title == "Carátula":
            continue
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.outlinePr.summaryBelow = True
        sheet.oddHeader.center.text = "&B UNIVERSIDAD NACIONAL DE INGENIERÍA - FIIS"
        sheet.oddFooter.right.text = "Página &P de &N"


def week_bounds(week: int) -> tuple[date, date]:
    start = date(2026, 9, 15) + timedelta(days=7 * (week - 1))
    return start, start + timedelta(days=6)


def make_gantt() -> Path:
    workbook = Workbook()
    add_cover(
        workbook,
        "CRONOGRAMA DE TRABAJO: 12 SEMANAS",
        "15-sep a 7-dic-2026. Cronograma auxiliar para revisar/completar el plan de tesis; no reemplaza el cronograma de investigación.",
    )
    sheet = workbook.create_sheet("Gantt")
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:S1")
    sheet["A1"] = "PLAN DE TRABAJO DEL PLAN DE TESIS | CORTE DE AVANCE: 10-OCT-2026"
    sheet["A1"].font = Font(name="Arial", bold=True, size=13, color=NAVY)
    sheet["A1"].alignment = Alignment(horizontal="center")
    sheet.merge_cells("A2:S2")
    sheet["A2"] = "S1-S3: productos/fuentes existentes sin fecha semanal verificable | S4: en curso | Semanas futuras: programadas"
    sheet["A2"].font = Font(name="Arial", italic=True, size=9, color=INK)
    sheet["A2"].alignment = Alignment(horizontal="center", wrap_text=True)
    headers = ["ID", "Actividad / producto", "Responsable", "Inicio", "Fin", "Depende de", "Estado"]
    headers.extend(f"S{week}\n{week_bounds(week)[0]:%d-%b}" for week in range(1, 13))
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=col, value=header)
        cell.fill = fill(NAVY)
        cell.font = Font(name="Arial", bold=True, size=9, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, (task_id, label, first, last, dependency, status) in enumerate(TASKS, start=5):
        start, _ = week_bounds(first)
        _, end = week_bounds(last)
        values = [task_id, label, "Tesistas" if task_id != 11 else "Tesistas + asesor", start, end, dependency, status]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col, value=value)
            cell.font = Font(name="Arial", size=9, color=INK)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color=GRID))
        sheet.cell(row_index, 4).number_format = "dd-mmm-yyyy"
        sheet.cell(row_index, 5).number_format = "dd-mmm-yyyy"
        for week in range(first, last + 1):
            cell = sheet.cell(row=row_index, column=7 + week, value="■")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
            if week <= 3 and task_id <= 4:
                cell.fill = fill(GREEN)
            elif week == 4 and task_id <= 5:
                cell.fill = fill(GOLD)
                cell.font = Font(name="Arial", size=10, bold=True, color=INK)
            else:
                cell.fill = fill(MID_BLUE)
        sheet.row_dimensions[row_index].height = 36
    legend_row = 19
    legends = [(GREEN, "Evidencia disponible al corte; fecha semanal no acreditada"), (GOLD, "En curso al 10-oct-2026"), (MID_BLUE, "Programado")]
    for offset, (color, text) in enumerate(legends):
        row = legend_row + offset
        sheet.cell(row, 1, "■").fill = fill(color)
        sheet.cell(row, 1).font = Font(name="Arial", color=WHITE, bold=True)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        sheet.cell(row, 2, text).font = Font(name="Arial", size=9, color=INK)
    sheet.merge_cells("I19:S19")
    sheet["I19"] = "Periodo: 15-sep–7-dic-2026 inclusive; semana 12 concluye el 7-dic."
    sheet["I19"].font = Font(name="Arial", size=9, italic=True, color=INK)
    widths = {"A": 5, "B": 48, "C": 19, "D": 13, "E": 13, "F": 13, "G": 43}
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width
    for col in range(8, 20):
        sheet.column_dimensions[chr(64 + col)].width = 10
    sheet.freeze_panes = "H5"
    sheet.auto_filter.ref = f"A4:S{4 + len(TASKS)}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.print_title_rows = "1:4"
    style_workbook(workbook)
    out = HERE / "Cronograma_12_semanas.xlsx"
    workbook.save(out)
    return out


def make_data_workbook(inventory: list[dict[str, str]], active_status: list[dict[str, object]], minem_raw: list[dict[str, object]]) -> Path:
    workbook = Workbook()
    add_cover(
        workbook,
        "INVENTARIO DE DATOS Y FUENTES",
        "Corte documental: 10-oct-2026 | incluye inventario previo, literatura/CIGRE y candidatos COES/MINEM",
    )

    case_rows = []
    for item in inventory:
        key = item["clave"]
        doi = DOI_FIXES.get(key, item.get("doi_o_identificador", ""))
        mapped = ACTIVE_MAP.get(key, "No reconciliado con caso activo")
        case_rows.append([
            key,
            item["fuente"],
            item["instancia"],
            item["tipo"],
            item["uso_pinn"],
            doi,
            mapped,
            "Resumen bibliográfico existente; validar campos y páginas fuente antes de modelar",
        ])
    sheet = workbook.create_sheet("Casos_16")
    write_table(sheet, ["Clave", "Fuente", "Instancia / descripción", "Tipo", "Uso previsto", "DOI / identificador cotejado", "Mapeo PINN_v2", "Nota de calidad"], case_rows, "CaseInventory")
    for row in range(2, sheet.max_row + 1):
        doi = sheet.cell(row, 6).value
        if isinstance(doi, str) and doi.startswith("10."):
            sheet.cell(row, 6).hyperlink = f"https://doi.org/{doi}"
            sheet.cell(row, 6).style = "Hyperlink"
    for col, width in {"A": 25, "B": 30, "C": 54, "D": 25, "E": 44, "F": 34, "G": 48, "H": 56}.items():
        sheet.column_dimensions[col].width = width

    coes_sheet = workbook.create_sheet("COES_cualquier_kV")
    coes_rows = [[r["id"], r["tipo"], r["codigo"], r["salida"], r["llegada"], r["kv"], r["km"], r["estado"], r["localizador"], r["datos"], r["fuente"]] for r in COES_ROWS]
    write_table(coes_sheet, ["ID", "Clase de evidencia", "Proyecto / circuito", "Origen", "Destino", "Tensión kV", "Longitud km", "Estado reportado", "Localizador", "Datos y límites", "Fuente local"], coes_rows, "COESCandidates")
    for col, width in {"A": 38, "B": 32, "C": 44, "D": 22, "E": 28, "F": 13, "G": 14, "H": 56, "I": 48, "J": 72, "K": 82}.items():
        coes_sheet.column_dimensions[col].width = width

    minem_sheet = workbook.create_sheet("MINEM_tramos")
    minem_rows = [[r["id"], r["codigo"], r["salida"], r["llegada"], r["kv"], r["km"], r["mva"], r["conductor"], r["mm2"], "Subterráneo", f"Texto extraído, línea {r['linea_txt']}", "Instalación en inventario MINEM 2024; no incluye geometría de zanja ni propiedades térmicas del suelo"] for r in MINEM_ROWS]
    write_table(minem_sheet, ["ID", "Código de línea", "SE salida", "SE llegada", "Tensión kV", "Longitud km", "Capacidad MVA", "Conductor", "Sección mm²", "Tipo", "Localizador en fuente", "Estado / faltantes relevantes"], minem_rows, "MINEMSegments")
    for col, width in {"A": 30, "B": 24, "C": 25, "D": 25, "E": 14, "F": 15, "G": 16, "H": 15, "I": 15, "J": 18, "K": 34, "L": 72}.items():
        minem_sheet.column_dimensions[col].width = width

    raw_sheet = workbook.create_sheet("MINEM_extracto_completo")
    raw_rows = [[row["id"], row["linea_txt"], row["codigo"], row["kv"], row["km"], row["mva"], row["conductor"], row["mm2"], row["parse_status"], row["raw"]] for row in minem_raw]
    write_table(raw_sheet, ["ID de extracto", "Línea del texto", "Código detectado", "kV leído", "Longitud leída km", "MVA leído", "Conductor leído", "Sección leída mm²", "Estado de parseo", "Texto extraído completo"], raw_rows, "MINEMAllMentions")
    for col, width in {"A": 34, "B": 16, "C": 26, "D": 14, "E": 18, "F": 14, "G": 28, "H": 18, "I": 48, "J": 120}.items():
        raw_sheet.column_dimensions[col].width = width

    field_rows = [
        ["Cable", "Material y capas; radios/diámetros; espesores; conductor", "m, mm, W/(m K)", "Geometría multicapa y ruta de calor; nunca completar capas no publicadas"],
        ["Instalación", "Coordenadas, formación, separación, profundidad, ductos, zanja, backfill y dominio", "m, mm", "Entrada geométrica principal; dato de activo MINEM/COES no suele bastar"],
        ["Suelo", "Estratos/capas, clasificación, k, resistividad, humedad, densidad y temperatura", "W/(m K), K m/W, %, kg/m³, °C", "Caracterizar k(x,y) y escenarios heterogéneos/estacionales"],
        ["Fronteras", "Temperatura superficial/ambiente, convección, frontera inferior/lateral", "°C, W/(m² K)", "Documentar condición y adaptación al dominio 2D"],
        ["Eléctrico/térmico", "Tensión, corriente, R20, factor de pérdidas, potencia/calor", "kV, A, Ω/m, W/m", "Conservar las magnitudes reportadas; tensión no determina calor generado"],
        ["Referencia", "Tmax, ampacidad, campo FEM/analítico, malla y convergencia", "°C, A", "Distinguir dato de referencia de salida PINN"],
        ["Procedencia", "Fuente, página/tabla/fila, fecha, estado y transformación", "texto + URI/DOI", "Estados PINN_v2: publicado, extraído, inferido, supuesto, faltante"],
    ]
    dictionary = workbook.create_sheet("Diccionario")
    write_table(dictionary, ["Grupo", "Campos requeridos", "Unidad esperada", "Relevancia / control"], field_rows, "FieldDictionary")
    for col, width in {"A": 22, "B": 76, "C": 37, "D": 72}.items():
        dictionary.column_dimensions[col].width = width

    source_rows = [
        ["Guía 1 UNI 2023", "Formato y estructura profesionalizante", "Plan/guias/GUÍA N°01 Plan.pdf; OCR local", "Estructura, tipografía y evaluación"],
        ["Informe de datos previo", "16 instancias y valores publicados", "Plan/envío_rev1/04_Datos_objeto_estudio/04_Informe_datos_casos_estudio_papers.tex; .md; 04_indice_instancias.csv", "Línea base; comprobar discrepancias"],
        ["CIGRE TB 880 (2022)", "Caso de cable 132 kV y variantes IEC", "Plan/envío_rev1/01_Referencias_Zotero/01_PDFs/; pp. 61-125 del informe previo", "Referencia normativa de ampacidad y pérdidas"],
        ["CIGRE TB 963 (2025)", "Caso cable FEM 2D y anillo analítico", "Plan/envío_rev1/01_Referencias_Zotero/01_PDFs/; caso 1 pp. 103-121; anillo pp. 52-55", "Campo y referencia exacta"],
        ["COES DP-01-2023", "Enlace 220 kV Chavarría-San Martín planificado; enlaces subterráneos 220 kV", "Plan/datos_fuente/extractos_texto/COES_-_2023_-_Informe_de_diagnstico_de_las_condiciones_operativas_del_SEIN__periodo_2025--2034.txt; proyecto p. 69", "Caso peruano; propuesta versus existente"],
        ["MINEM Anuario 2024", "Tramos subterráneos con código, ruta, longitud, tensión, capacidad y conductor", "Plan/datos_fuente/descargas/minem_anuario_2024.pdf; texto extraído, tabla de líneas", "Evidencia de activos, geometría térmica incompleta"],
        ["Arias Velasquez et al. 2023", "Diseño térmico de circuitos subterráneos MV", "DOI 10.1016/j.rineng.2023.101552", "Artículo parecido; no sustituye fuente de activo local"],
        ["PINN_v2 activo", "21 definiciones JSON vigentes", "PINN_v2/pinn_v2/cases/v2/; schema pinn-v2-case-1", "Validación de casos ejecutables y estado de procedencia"],
    ]
    sources = workbook.create_sheet("Fuentes")
    write_table(sources, ["Fuente / identificador", "Contenido", "Localizador / archivo", "Aporte y límite"], source_rows, "SourceRegister")
    for col, width in {"A": 32, "B": 58, "C": 100, "D": 60}.items():
        sources.column_dimensions[col].width = width

    preprocessing_rows = [
        [1, "Inventario", "Preservar las 16 claves del CSV previo; agregar IDs estables a COES/MINEM", "Sin deduplicar circuitos distintos; repetir código/longitud se marca como fila separada"],
        [2, "Transcripción", "Conservar unidad/valor original, código y localizador de texto/página", "No ajustar cifras por consistencia con otro artículo"],
        [3, "Normalización", "Convertir solo al exportar campos del modelo; mantener valor original", "No convertir MVA en corriente/calor sin tensión de operación y condiciones"],
        [4, "Procedencia", "Asignar publicado/extraído/inferido/supuesto/faltante por grupo", "Un JSON que carga puede seguir teniendo valores supuestos"],
        [5, "Reconciliación", "Cotejar DOIs con Plan/referencias_plan_tesis.bib y títulos", "Corregir discrepancias del CSV anterior; p.ej., Atoccsa DOI 10.3390/en17051023"],
        [6, "Modelabilidad", "Separar inventario/proyecto de caso computable", "No inferir zanja, profundidad, arreglo ni estratos de la tensión"],
        [7, "PINN_v2", "Copiar JSON v2 activo; validar con CaseDefinition.load", "Candidatos no mapeados quedan fuera del árbol ejecutable"],
    ]
    prep = workbook.create_sheet("Preprocesamiento")
    write_table(prep, ["Paso", "Etapa", "Operación", "Regla / límite"], preprocessing_rows, "Preprocessing")
    for col, width in {"A": 10, "B": 22, "C": 78, "D": 78}.items():
        prep.column_dimensions[col].width = width

    quality_rows = []
    for row in COES_ROWS:
        quality_rows.append([row["id"], row["tipo"], "geometry; materials; source heat; boundary_conditions", "No; candidato", "Fuente primaria verificable, pero faltan sección de cable/instalación y propiedades del terreno"])
    for row in minem_raw:
        quality_rows.append([row["id"], "Mención MINEM 2024 no deduplicada", "geometry; materials; soil; boundary_conditions; heat", "No; extracto de inventario", "El renglón puede registrar un tramo parcial o duplicado; faltan dimensiones de capas, formación, profundidad y suelo térmico"])
    quality = workbook.create_sheet("Calidad_faltantes")
    write_table(quality, ["ID", "Clase de evidencia", "Grupos v2 pendientes", "¿Caso modelable?", "Faltantes principales"], quality_rows, "DataQuality")
    for col, width in {"A": 42, "B": 42, "C": 62, "D": 25, "E": 82}.items():
        quality.column_dimensions[col].width = width

    json_sheet = workbook.create_sheet("PINN_v2_activos")
    json_rows = [[r["id"], r["tipo"], r["estado_cargador"], r["grupos_publicado"], r["grupos_extraido"], r["grupos_supuesto"], r["procedencia"]] for r in active_status]
    write_table(json_sheet, ["ID activo", "Kind", "Validación", "Grupos publicado", "Grupos extraído", "Grupos supuesto", "Estado por grupo"], json_rows, "ActiveV2Cases")
    for col, width in {"A": 30, "B": 22, "C": 34, "D": 18, "E": 18, "F": 18, "G": 86}.items():
        json_sheet.column_dimensions[col].width = width

    style_workbook(workbook)
    workbook.properties.title = "Inventario de datos y fuentes para tesis profesionalizante"
    workbook.properties.creator = AUTHORS
    out = HERE / "Inventario_datos.xlsx"
    workbook.save(out)
    return out


def candidate_case(case_id: str, title: str, source_id: str, citation: str, reported_data: dict[str, object]) -> dict[str, object]:
    groups = ("geometry", "materials", "source", "boundary_conditions")
    return {
        "schema_version": SCHEMA_VERSION,
        "id": case_id,
        "title": title,
        "physical": {"id": case_id},
        "boundary_conditions": None,
        "contact": [],
        "provenance": [{"id": source_id, "citation": citation, "reported_data": reported_data}],
        "data_status": {
            group: {
                "status": "faltante",
                "source": source_id,
                "note": "No se ha normalizado el grupo requerido al bloque físico PINN_v2; registro de candidato, no ejecutable.",
            }
            for group in groups
        },
        "inconsistencies": [],
    }


def make_json_zip(inventory: list[dict[str, str]], active: list[tuple[Path, CaseDefinition]], minem_raw: list[dict[str, object]]) -> Path:
    archive_path = HERE / "Datos_casos_PINN_v2_20261010.zip"
    non_active_inventory = [item for item in inventory if item["clave"] not in ACTIVE_MAP]
    manifest: dict[str, object] = {
        "schema_version": SCHEMA_VERSION,
        "package_date": "2026-10-10",
        "executable_case_count": len(active),
        "candidate_only_count": len(non_active_inventory) + len(COES_ROWS) + len(minem_raw),
        "executable_cases": [case.id for _, case in active],
        "candidate_inventory": "candidatos_no_ejecutables/indice_candidatos.json",
        "validation": "Each active v2 case was loaded through CaseDefinition.load during generation.",
        "candidate_note": "Candidate records preserve source-reported metadata but are intentionally rejected by the current solver loader because required model fields are not mapped or are missing. They do not imply missing information in the cited paper when the inventory summary has partial values.",
    }
    with zipfile.ZipFile(archive_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path, case in active:
            relative = path.relative_to(PINN_ROOT / "pinn_v2" / "cases" / "v2")
            archive.write(path, f"ejecutables/{relative.as_posix()}")
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
        index_records = []
        for item in inventory:
            key = item["clave"]
            doi = DOI_FIXES.get(key, item.get("doi_o_identificador", ""))
            index_records.append(
                {
                    "id": key,
                    "title": item["instancia"],
                    "record_type": item["tipo"],
                    "source": item["fuente"],
                    "reported_use": item["uso_pinn"],
                    "doi_or_identifier": doi,
                    "active_v2_match": ACTIVE_MAP.get(key),
                    "status": "candidate_only" if key not in ACTIVE_MAP else "see_executable_case",
                }
            )
        for row in COES_ROWS:
            index_records.append(
                {
                    "id": row["id"],
                    "title": row["codigo"],
                    "record_type": row["tipo"],
                    "source": "COES/DP-01-2023",
                    "reported_use": "Peruvian installation/project evidence; not model-ready",
                    "doi_or_identifier": row["localizador"],
                    "status": "candidate_only",
                }
            )
        for row in minem_raw:
            index_records.append(
                {
                    "id": row["id"],
                    "title": f"MINEM 2024 underground-row extract {row['linea_txt']}",
                    "record_type": "Peruvian inventory text mention; not deduplicated",
                    "source": "MINEM Anuario Estadístico de Electricidad 2024",
                    "reported_use": "Evidence of underground transmission segment; not model-ready",
                    "doi_or_identifier": row["codigo"],
                    "source_line": row["linea_txt"],
                    "status": "candidate_only",
                }
            )
        archive.writestr(
            "candidatos_no_ejecutables/indice_candidatos.json",
            json.dumps({"schema_version": SCHEMA_VERSION, "count": len(index_records), "cases": index_records}, ensure_ascii=False, indent=2) + "\n",
        )
        for item in inventory:
            key = item["clave"]
            if key in ACTIVE_MAP:
                continue
            doi = DOI_FIXES.get(key, item.get("doi_o_identificador", ""))
            candidate = candidate_case(
                key,
                item["instancia"],
                "prior_inventory",
                f"04_indice_instancias.csv; source {item['fuente']}; DOI/ID {doi}",
                {"record_type": item["tipo"], "summary": item["instancia"], "intended_use": item["uso_pinn"], "doi_or_identifier": doi},
            )
            archive.writestr(f"candidatos_no_ejecutables/{key}.json", json.dumps(candidate, ensure_ascii=False, indent=2) + "\n")
        for row in COES_ROWS:
            candidate = candidate_case(
                row["id"], row["codigo"], "coes2023diagnostico",
                f"COES/DP-01-2023, {row['localizador']}; extracted text in {row['fuente']}",
                {key: row[key] for key in ("tipo", "codigo", "salida", "llegada", "kv", "km", "estado", "datos")},
            )
            archive.writestr(f"candidatos_no_ejecutables/{row['id']}.json", json.dumps(candidate, ensure_ascii=False, indent=2) + "\n")
        for row in minem_raw:
            candidate = candidate_case(
                row["id"], f"MINEM 2024 underground row {row['linea_txt']}", "minem_anuario_2024",
                f"MINEM, Anuario Estadístico de Electricidad 2024, texto extraído, línea {row['linea_txt']}",
                {key: row[key] for key in ("codigo", "kv", "km", "mva", "conductor", "mm2", "linea_txt", "raw", "parse_status")},
            )
            archive.writestr(f"candidatos_no_ejecutables/{row['id']}.json", json.dumps(candidate, ensure_ascii=False, indent=2) + "\n")
        readme = (
            "# Paquete de casos y candidatos PINN_v2\n\n"
            f"Versión de esquema: `{SCHEMA_VERSION}`. Fecha del inventario: 2026-10-10.\n\n"
            f"El ZIP incluye {len(active)} JSON activos copiados desde `PINN_v2/pinn_v2/cases/v2/`; cada uno fue cargado con `CaseDefinition.load` al generar el paquete. "
            "La hoja de cálculo asociada informa la procedencia por grupo: que una definición cargue no significa que todos sus valores sean publicados; varios registros heredados mantienen `supuesto`.\n\n"
            f"La carpeta `candidatos_no_ejecutables/` contiene el índice de las 16 entradas del informe previo, registros COES y {len(minem_raw)} menciones textuales MINEM de `Subterráneo` sin deduplicar. Estos JSON reutilizan la envoltura y el `schema_version` de PINN_v2, pero dejan los grupos requeridos como `faltante`; por diseño no pasan el cargador ni deben usarse para entrenar/simular. Para varios artículos, `reported_data` resume lo que la fuente sí publica; `faltante` significa que no está normalizado en el bloque físico completo, no necesariamente que el artículo carezca de ese dato. No se inventan dimensiones, profundidad, capas de suelo ni fronteras.\n\n"
            "El proyecto COES de 220 kV se clasifica como planificado según el diagnóstico 2023. Los registros MINEM son inventario de líneas y no contienen geometría de zanja ni propiedades térmicas del suelo. El PDF separado COES/DP-02-2022 no pudo descargarse desde la URL citada (respuesta vacía); el caso COES se sustenta en el diagnóstico local consultado.\n"
        )
        archive.writestr("README.md", readme)
    return archive_path


def validate_outputs(inventory: list[dict[str, str]], active_count: int, minem_raw: list[dict[str, object]]) -> None:
    inventory_count = len(inventory)
    data_path = HERE / "Inventario_datos.xlsx"
    gantt_path = HERE / "Cronograma_12_semanas.xlsx"
    archive_path = HERE / "Datos_casos_PINN_v2_20261010.zip"
    data_book = load_workbook(data_path, read_only=True, data_only=True)
    assert {"Carátula", "Casos_16", "COES_cualquier_kV", "MINEM_tramos", "MINEM_extracto_completo", "Diccionario", "Fuentes", "Preprocesamiento", "Calidad_faltantes", "PINN_v2_activos"} <= set(data_book.sheetnames)
    assert data_book["Casos_16"].max_row == inventory_count + 1
    assert data_book["COES_cualquier_kV"].max_row == len(COES_ROWS) + 1
    assert data_book["MINEM_tramos"].max_row == len(MINEM_ROWS) + 1
    assert data_book["MINEM_extracto_completo"].max_row == len(minem_raw) + 1
    assert data_book["PINN_v2_activos"].max_row == active_count + 1
    gantt_book = load_workbook(gantt_path, read_only=True, data_only=True)
    gantt = gantt_book["Gantt"]
    assert gantt.max_row >= 16 and gantt.max_column == 19
    assert [gantt.cell(row=4, column=column).value.split("\n")[0] for column in range(8, 20)] == [f"S{week}" for week in range(1, 13)]
    with zipfile.ZipFile(archive_path) as archive:
        names = archive.namelist()
        active_names = [name for name in names if name.startswith("ejecutables/") and name.endswith(".json")]
        candidate_names = [name for name in names if name.startswith("candidatos_no_ejecutables/") and name.endswith(".json") and not name.endswith("indice_candidatos.json")]
        manifest = json.loads(archive.read("manifest.json"))
        assert len(active_names) == active_count
        assert len(candidate_names) == manifest["candidate_only_count"]
        assert len(candidate_names) == len([item for item in inventory if item["clave"] not in ACTIVE_MAP]) + len(COES_ROWS) + len(minem_raw)
        for name in active_names:
            CaseDefinition.from_dict(json.loads(archive.read(name)))
        for name in candidate_names:
            candidate = json.loads(archive.read(name))
            assert candidate["schema_version"] == SCHEMA_VERSION
            assert all(candidate["data_status"][group]["status"] == "faltante" for group in ("geometry", "materials", "source", "boundary_conditions"))
            try:
                CaseDefinition.from_dict(candidate)
            except Exception:
                continue
            raise AssertionError(f"Candidate unexpectedly accepted by PINN_v2 loader: {name}")
    print(f"VALIDATED workbooks: prior={inventory_count}, COES={len(COES_ROWS)}, MINEM_examples={len(MINEM_ROWS)}, MINEM_extracts={len(minem_raw)}, active={active_count}")
    print(f"VALIDATED ZIP: executable={len(active_names)}, non_executable_candidates={len(candidate_names)}")


def main() -> None:
    inventory = read_prior_inventory()
    minem_raw = read_minem_extracts()
    active, active_status = load_active_cases()
    if len(inventory) != 16:
        raise RuntimeError(f"Expected 16 prior inventory rows, found {len(inventory)}")
    if len(active) != 21:
        raise RuntimeError(f"Expected 21 active v2 cases, found {len(active)}")
    if len(minem_raw) != 139:
        raise RuntimeError(f"Expected 139 raw MINEM underground mentions, found {len(minem_raw)}")
    outputs = [make_gantt(), make_data_workbook(inventory, active_status, minem_raw), make_json_zip(inventory, active, minem_raw)]
    validate_outputs(inventory, len(active), minem_raw)
    for output in outputs:
        print(f"CREATED {output.relative_to(HERE)} ({output.stat().st_size} bytes)")
    print(f"CHECKED prior_inventory={len(inventory)} active_v2={len(active)} minem_examples={len(MINEM_ROWS)} minem_extracts={len(minem_raw)} coes_records={len(COES_ROWS)}")


if __name__ == "__main__":
    main()