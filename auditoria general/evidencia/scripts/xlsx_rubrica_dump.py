import openpyxl, sys, json
p = r'c:\usr\ths_mia_fiis\tesis_MIA_lk\Plan\evaluacion_formato_8_2_plan_tesis.xlsx'
wb = openpyxl.load_workbook(p)
out = {}
for ws in wb.worksheets:
    print('=== HOJA', ws.title, ws.max_row, ws.max_column, 'merged:', len(ws.merged_cells.ranges))
    for r in ws.iter_rows():
        vals = [(c.coordinate, c.value) for c in r if c.value is not None]
        if vals:
            print(' | '.join(f'{k}={str(v)[:400]!r}' for k, v in vals))
print('col widths', {ws.title: {k: v.width for k, v in ws.column_dimensions.items()} for ws in wb.worksheets})
