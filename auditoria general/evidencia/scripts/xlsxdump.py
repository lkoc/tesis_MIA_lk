import os, sys
os.chdir(r'c:\usr\ths_mia_fiis\tesis_MIA_lk')
try:
    import openpyxl
except ImportError:
    print('no openpyxl'); sys.exit()
for p in ['matriz_consistencia/Matriz_de_Consistencia_Tesis_Cables_Heterogeneidad_Termica.xlsx', 'Plan/evaluacion_formato_8_2_plan_tesis.xlsx']:
    wb = openpyxl.load_workbook(p, data_only=True)
    print('=====', p, wb.sheetnames, 'mtime', os.path.getmtime(p))
    for ws in wb.worksheets:
        print('--- hoja', ws.title, ws.max_row, 'x', ws.max_column)
        n = 0
        for row in ws.iter_rows(values_only=True):
            vals = [str(v).strip().replace('\n', ' ')[:160] for v in row if v is not None and str(v).strip()]
            if vals:
                print('  |', ' | '.join(vals)[:600]); n += 1
            if n >= 45: break
